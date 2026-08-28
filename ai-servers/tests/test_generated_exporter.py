import hashlib
import base64
import json
import os
import tempfile
import unittest
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from docx import Document

from app.rag.document_conversion import generated_exporter


class GeneratedExporterTest(unittest.TestCase):
    def setUp(self):
        self.original_root = generated_exporter.EXPORT_ROOT
        self.original_ttl_hours = getattr(generated_exporter, "EXPORT_TTL_HOURS", None)
        self.original_max_bytes = getattr(generated_exporter, "EXPORT_MAX_BYTES", None)

    def tearDown(self):
        generated_exporter.EXPORT_ROOT = self.original_root
        if self.original_ttl_hours is not None:
            generated_exporter.EXPORT_TTL_HOURS = self.original_ttl_hours
        if self.original_max_bytes is not None:
            generated_exporter.EXPORT_MAX_BYTES = self.original_max_bytes

    @staticmethod
    def _markdown_metadata():
        return {
            "executedAgent": "textbook_knowledge_agent",
            "toolToggles": {
                "docx_export_tool": False,
                "excel_export_tool": False,
                "content_archive_tool": False,
            },
        }

    @staticmethod
    def _read_manifest(root: Path, storage_key: str):
        return json.loads((root / f"{storage_key}.meta.json").read_text(encoding="utf-8"))

    def test_generated_image_answer_becomes_private_export_without_provider_payload(self):
        image_bytes = b"\x89PNG\r\n\x1a\n" + b"private-image"
        answer = json.dumps({
            "status": "success",
            "message": "generated",
            "images": [{
                "index": 0,
                "url": "https://temporary.example/image.png",
                "base64": base64.b64encode(image_bytes).decode("ascii"),
                "contentType": "image/png",
                "status": "success",
            }],
        })
        with tempfile.TemporaryDirectory() as directory:
            generated_exporter.EXPORT_ROOT = Path(directory)
            clean_answer, attachments = generated_exporter.materialize_generated_image_answer(
                answer,
                display_stem="Python 学习流程图",
                tool_name="generate_flowchart_image_tool",
            )

            clean_payload = json.loads(clean_answer)
            self.assertNotIn("url", clean_payload["images"][0])
            self.assertNotIn("base64", clean_payload["images"][0])
            self.assertEqual("image", attachments[0]["type"])
            self.assertEqual("Python 学习流程图.png", attachments[0]["fileName"])
            self.assertTrue(attachments[0]["serverGenerated"])
            self.assertEqual(
                image_bytes,
                (Path(directory) / attachments[0]["storageKey"]).read_bytes(),
            )

    def test_docx_export_embeds_workflow_generated_image(self):
        image_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory() as directory:
            generated_exporter.EXPORT_ROOT = Path(directory)
            result = generated_exporter.export_generated_answer(
                "# 测试文档\n\n这是正文。",
                "markdown",
                {
                    "requestedOutputType": "docx",
                    "embeddedImageBytes": [image_bytes],
                    "toolToggles": {"docx_export_tool": True},
                },
            )

            document = Document(Path(directory) / result.attachments[0]["storageKey"])
            self.assertEqual(1, len(document.inline_shapes))

    def test_default_root_is_repository_local_and_production_requires_explicit_root(self):
        expected = Path(generated_exporter.__file__).resolve().parents[3] / "data" / "ai-exports"

        self.assertEqual(expected.resolve(), generated_exporter._resolve_export_root({}))
        with self.assertRaisesRegex(RuntimeError, "AI_EXPORT_ROOT"):
            generated_exporter._resolve_export_root({"AI_ENV": "production"})
        self.assertEqual(
            Path("/srv/shared/ai-exports"),
            generated_exporter._resolve_export_root({
                "AI_ENV": "prod",
                "AI_EXPORT_ROOT": "/srv/shared/ai-exports",
            }),
        )

    def test_export_uses_full_uuid_storage_key_and_atomic_digest_sidecar(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            generated_exporter.EXPORT_ROOT = root
            generated_exporter.EXPORT_TTL_HOURS = 168
            generated_exporter.EXPORT_MAX_BYTES = 1024 * 1024
            destinations = []
            real_replace = generated_exporter.os.replace

            def record_replace(source, destination):
                destinations.append(Path(destination).name)
                return real_replace(source, destination)

            with patch.object(generated_exporter.os, "replace", side_effect=record_replace):
                result = generated_exporter.export_generated_answer(
                    "# 栈与队列\n\n- 栈：后进先出",
                    "markdown",
                    self._markdown_metadata(),
                )

            self.assertEqual(1, len(result.attachments))
            attachment = result.attachments[0]
            required_fields = {
                "storageKey",
                "serverGenerated",
                "internalCapability",
                "sha256",
                "size",
                "createdAt",
                "expiresAt",
            }
            self.assertTrue(required_fields.issubset(attachment))
            self.assertTrue(attachment["serverGenerated"])
            self.assertNotIn("url", attachment)
            storage_key = attachment["storageKey"]
            self.assertTrue(attachment["name"].endswith(".md"))
            self.assertNotEqual("Markdown.md", attachment["name"])
            self.assertEqual(attachment["name"], attachment["fileName"])
            self.assertNotEqual(storage_key, attachment["name"])
            self.assertEqual(str(uuid.UUID(Path(storage_key).stem)), Path(storage_key).stem)

            payload_path = root / storage_key
            sidecar_path = root / f"{storage_key}.meta.json"
            self.assertTrue(payload_path.is_file())
            self.assertTrue(sidecar_path.is_file())
            self.assertLess(destinations.index(storage_key), destinations.index(sidecar_path.name))
            self.assertFalse(any(".tmp" in item.name for item in root.iterdir()))

            manifest = self._read_manifest(root, storage_key)
            self.assertEqual(
                {"capabilityDigest", "sha256", "size", "mimeType", "createdAt", "expiresAt"},
                set(manifest),
            )
            self.assertNotIn(attachment["internalCapability"], sidecar_path.read_text(encoding="utf-8"))
            self.assertEqual(
                hashlib.sha256(attachment["internalCapability"].encode("utf-8")).hexdigest(),
                manifest["capabilityDigest"],
            )
            self.assertEqual(hashlib.sha256(payload_path.read_bytes()).hexdigest(), manifest["sha256"])
            self.assertEqual(payload_path.stat().st_size, manifest["size"])
            self.assertEqual(manifest["sha256"], attachment["sha256"])
            self.assertEqual(manifest["size"], attachment["size"])
            created_at = datetime.fromisoformat(manifest["createdAt"].replace("Z", "+00:00"))
            expires_at = datetime.fromisoformat(manifest["expiresAt"].replace("Z", "+00:00"))
            self.assertEqual(timedelta(hours=168), expires_at - created_at)

    def test_cleanup_removes_orphans_expired_pairs_and_oldest_pair_over_capacity(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            generated_exporter.EXPORT_ROOT = root
            generated_exporter.EXPORT_TTL_HOURS = 168
            generated_exporter.EXPORT_MAX_BYTES = 1024 * 1024

            first = generated_exporter.export_generated_answer(
                "# 第一份\n\n- 旧内容",
                "markdown",
                self._markdown_metadata(),
            ).attachments[0]
            second = generated_exporter.export_generated_answer(
                "# 第二份\n\n- 新内容更多一些",
                "markdown",
                self._markdown_metadata(),
            ).attachments[0]

            now = datetime.now(timezone.utc)
            first_manifest = self._read_manifest(root, first["storageKey"])
            first_manifest["createdAt"] = (now - timedelta(hours=2)).isoformat().replace("+00:00", "Z")
            first_manifest["expiresAt"] = (now - timedelta(hours=1)).isoformat().replace("+00:00", "Z")
            (root / f"{first['storageKey']}.meta.json").write_text(
                json.dumps(first_manifest), encoding="utf-8"
            )

            orphan_payload = root / f"{uuid.uuid4()}.txt"
            orphan_payload.write_text("orphan", encoding="utf-8")
            orphan_sidecar = root / f"{uuid.uuid4()}.txt.meta.json"
            orphan_sidecar.write_text("{}", encoding="utf-8")
            stale_orphan_time = (now - timedelta(seconds=301)).timestamp()
            os.utime(orphan_payload, (stale_orphan_time, stale_orphan_time))
            os.utime(orphan_sidecar, (stale_orphan_time, stale_orphan_time))

            generated_exporter.cleanup_generated_exports(
                root=root,
                now=now,
                max_bytes=1024 * 1024,
            )

            self.assertFalse((root / first["storageKey"]).exists())
            self.assertFalse((root / f"{first['storageKey']}.meta.json").exists())
            self.assertFalse(orphan_payload.exists())
            self.assertFalse(orphan_sidecar.exists())
            self.assertTrue((root / second["storageKey"]).exists())

            third = generated_exporter.export_generated_answer(
                "# 第三份\n\n- 最新内容",
                "markdown",
                self._markdown_metadata(),
            ).attachments[0]
            capacity = third["size"]
            generated_exporter.cleanup_generated_exports(
                root=root,
                now=now,
                max_bytes=capacity,
            )

            self.assertFalse((root / second["storageKey"]).exists())
            self.assertFalse((root / f"{second['storageKey']}.meta.json").exists())
            self.assertTrue((root / third["storageKey"]).exists())
            self.assertTrue((root / f"{third['storageKey']}.meta.json").exists())

    def test_cleanup_preserves_fresh_staging_and_only_removes_stale_orphans(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            generated_exporter.EXPORT_ROOT = root
            generated_exporter.EXPORT_TTL_HOURS = 168
            generated_exporter.EXPORT_MAX_BYTES = 1024 * 1024
            now = datetime.now(timezone.utc)
            stale_time = (now - timedelta(seconds=301)).timestamp()
            fresh_time = (now - timedelta(seconds=299)).timestamp()

            fresh_temp = root / ".fresh-payload.tmp"
            stale_temp = root / ".stale-payload.tmp"
            fresh_payload = root / f"{uuid.uuid4()}.md"
            stale_payload = root / f"{uuid.uuid4()}.md"
            for path, content, modified_at in (
                (fresh_temp, "fresh-temp", fresh_time),
                (stale_temp, "stale-temp", stale_time),
                (fresh_payload, "fresh-payload", fresh_time),
                (stale_payload, "stale-payload", stale_time),
            ):
                path.write_text(content, encoding="utf-8")
                os.utime(path, (modified_at, modified_at))

            generated_exporter.cleanup_generated_exports(
                root=root,
                now=now,
                max_bytes=1024 * 1024,
                staging_grace_seconds=300,
            )

            self.assertTrue(fresh_temp.exists())
            self.assertTrue(fresh_payload.exists())
            self.assertFalse(stale_temp.exists())
            self.assertFalse(stale_payload.exists())

    def test_capacity_counts_only_committed_payload_sidecar_pairs(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            generated_exporter.EXPORT_ROOT = root
            generated_exporter.EXPORT_TTL_HOURS = 168
            generated_exporter.EXPORT_MAX_BYTES = 1024 * 1024
            committed = generated_exporter.export_generated_answer(
                "# 已提交\n\n- 内容",
                "markdown",
                self._markdown_metadata(),
            ).attachments[0]
            fresh_unpaired = root / f"{uuid.uuid4()}.md"
            fresh_unpaired.write_bytes(b"x" * (committed["size"] + 4096))

            generated_exporter.cleanup_generated_exports(
                root=root,
                now=datetime.now(timezone.utc),
                max_bytes=committed["size"],
                staging_grace_seconds=300,
            )

            self.assertTrue((root / committed["storageKey"]).exists())
            self.assertTrue((root / f"{committed['storageKey']}.meta.json").exists())
            self.assertTrue(fresh_unpaired.exists())

    def test_question_bank_exports_markdown_docx_and_xlsx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generated_exporter.EXPORT_ROOT = Path(temp_dir)
            payload = {
                "questions": [
                    {
                        "id": "SC1",
                        "type": "single_choice",
                        "stem": "栈的典型访问规则是什么？",
                        "score": 2,
                        "difficulty": "easy",
                        "knowledgePoints": ["栈"],
                        "tags": ["数据结构"],
                        "body": {
                            "options": [
                                {"key": "A", "text": "后进先出"},
                                {"key": "B", "text": "先进先出"},
                            ],
                        },
                        "answer": {"correctOption": "A"},
                        "analysis": "栈遵循后进先出。",
                        "scoring": {"mode": "exact", "rubrics": [{"criterion": "答对", "score": 2}]},
                        "sourceBasis": ["教材"],
                    },
                ],
                "missingInfo": [],
            }

            result = generated_exporter.export_generated_answer(
                json.dumps(payload, ensure_ascii=False),
                "question_bank",
                {"executedAgent": "textbook_question_single_choice_agent"},
            )

            self.assertEqual(["md", "docx", "xlsx", "zip"], [item["ext"] for item in result.attachments])
            self.assertEqual("question_bank", result.diagnostics["contentKind"])
            for attachment in result.attachments:
                path = generated_exporter.EXPORT_ROOT / attachment["storageKey"]
                self.assertTrue(path.exists(), path)
            xlsx_key = next(item["storageKey"] for item in result.attachments if item["ext"] == "xlsx")
            with zipfile.ZipFile(generated_exporter.EXPORT_ROOT / xlsx_key) as archive:
                self.assertIn("xl/worksheets/sheet1.xml", archive.namelist())

    def test_question_bank_export_tolerates_fsync_failures(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generated_exporter.EXPORT_ROOT = Path(temp_dir)
            payload = {
                "questions": [
                    {
                        "id": "SC1",
                        "type": "single_choice",
                        "stem": "什么结构满足后进先出？",
                        "body": {
                            "options": [
                                {"key": "A", "text": "栈"},
                                {"key": "B", "text": "队列"},
                            ],
                        },
                        "answer": {"correctOption": "A"},
                    },
                ],
                "missingInfo": [],
            }

            with patch.object(generated_exporter.os, "fsync", side_effect=OSError(9, "Bad file descriptor")):
                result = generated_exporter.export_generated_answer(
                    json.dumps(payload, ensure_ascii=False),
                    "question_bank",
                    {"executedAgent": "textbook_question_single_choice_agent"},
                )

            self.assertEqual(["md", "docx", "xlsx", "zip"], [item["ext"] for item in result.attachments])
            for attachment in result.attachments:
                self.assertTrue((generated_exporter.EXPORT_ROOT / attachment["storageKey"]).exists())

    def test_markdown_exports_reading_files(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generated_exporter.EXPORT_ROOT = Path(temp_dir)
            result = generated_exporter.export_generated_answer(
                "# 栈与队列\n\n- 栈：后进先出\n- 队列：先进先出",
                "markdown",
                {"executedAgent": "textbook_knowledge_agent"},
            )

            self.assertEqual(["md", "docx", "xlsx", "zip"], [item["ext"] for item in result.attachments])
            self.assertEqual("markdown_content", result.diagnostics["contentKind"])

    def test_requested_file_format_only_returns_selected_attachment(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generated_exporter.EXPORT_ROOT = Path(temp_dir)
            content = "# 栈与队列\n\n- 栈：后进先出\n- 队列：先进先出"

            word_result = generated_exporter.export_generated_answer(
                content,
                "markdown",
                {
                    "executedAgent": "generated_export_tools",
                    "allowGeneratedExportTool": True,
                    "requestedOutputType": "docx",
                },
            )
            excel_result = generated_exporter.export_generated_answer(
                content,
                "markdown",
                {
                    "executedAgent": "generated_export_tools",
                    "allowGeneratedExportTool": True,
                    "requestedOutputType": "xlsx",
                },
            )
            ppt_result = generated_exporter.export_generated_answer(
                "# 栈与队列\n\n## 核心概念\n\n- 栈：后进先出\n- 队列：先进先出",
                "markdown",
                {
                    "executedAgent": "generated_export_tools",
                    "allowGeneratedExportTool": True,
                    "requestedOutputType": "pptx",
                },
            )

            self.assertEqual(["docx"], [item["ext"] for item in word_result.attachments])
            self.assertEqual(["xlsx"], [item["ext"] for item in excel_result.attachments])
            self.assertEqual(["pptx"], [item["ext"] for item in ppt_result.attachments])
            self.assertTrue(word_result.attachments[0]["name"].endswith(".docx"))
            self.assertTrue(excel_result.attachments[0]["name"].endswith(".xlsx"))
            self.assertTrue(ppt_result.attachments[0]["name"].endswith(".pptx"))
            self.assertNotEqual("Word 文档.docx", word_result.attachments[0]["name"])
            self.assertNotEqual("Excel 表格.xlsx", excel_result.attachments[0]["name"])
            presentation = Presentation(
                generated_exporter.EXPORT_ROOT / ppt_result.attachments[0]["storageKey"]
            )
            self.assertGreaterEqual(len(presentation.slides), 2)

    def test_mermaid_exports_source_files(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            generated_exporter.EXPORT_ROOT = Path(temp_dir)
            result = generated_exporter.export_generated_answer(
                "```mermaid\nflowchart TD\n  A[开始] --> B[结束]\n```",
                "mermaid_flowchart",
                {"executedAgent": "diagram_flowchart_agent"},
            )

            self.assertEqual(["mmd", "md", "zip"], [item["ext"] for item in result.attachments])
            self.assertEqual("diagram_source", result.diagnostics["contentKind"])
            mmd_key = next(item["storageKey"] for item in result.attachments if item["ext"] == "mmd")
            self.assertIn("flowchart TD", (generated_exporter.EXPORT_ROOT / mmd_key).read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
